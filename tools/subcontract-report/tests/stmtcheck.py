# Cost Statement: the cached spill values must equal the statement recomputed from the recalculated Detail (DIV (live), Statement key, bucket, cost)
import openpyxl, subprocess, tempfile, shutil, sys, warnings; warnings.filterwarnings('ignore')
src = sys.argv[1]; d = tempfile.mkdtemp(); shutil.copy(src, d + '/b.xlsx')
subprocess.run(['soffice', '--headless', '--norestore', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', d + '/o', d + '/b.xlsx'], capture_output=True, timeout=900)
calc = openpyxl.load_workbook(d + '/o/b.xlsx', data_only=True); raw = openpyxl.load_workbook(src, data_only=True)
dt = calc['Detail']; H = [c.value for c in dt[1]]; ix = {h: i for i, h in enumerate(H)}
cs = raw['Cost Statement']; b3 = cs['B3'].value; M = b3.year * 100 + b3.month; PM = M - 89 if M % 100 == 1 else M - 1; YS = M // 100 * 100 + 1
def agg(keycol):
    out = {}
    for r in dt.iter_rows(min_row=2, values_only=True):
        k = r[ix['Statement key']]
        if not k: continue
        k = r[ix[keycol]]; b = r[ix['Report bucket']]; a = r[ix['Cost (excl VAT)']] or 0; x = out.setdefault(k, [0] * 5)
        if b == M: x[0] += a
        if b == PM: x[1] += a
        if isinstance(b, (int, float)) and YS <= b <= M: x[2] += a
        if (isinstance(b, (int, float)) and 0 < b <= M) or b == 'OPENING': x[3] += a
        if b == 'PENDING': x[4] += a
    return out
bad = 0
for name, keycol, r0 in [('DIV', 'DIV (live)', 7), ('service', 'Statement key', None)]:
    if r0 is None: r0 = next(r for r in range(8, 400) if cs.cell(r, 1).value == 'Key') + 1
    exp = agg(keycol); keys = sorted(exp, key=str.upper); got = []
    r = r0
    while cs.cell(r, 1).value not in (None, ''): got.append(cs.cell(r, 1).value); r += 1
    if got != keys: print(name, 'row list differs', got[:5], keys[:5]); bad += 1
    for n, k in enumerate(keys):
        row = [cs.cell(r0 + n, c).value for c in (6, 7, 10, 11, 12)]; e = exp[k]
        want = [e[0], e[1], e[2], e[3], e[4]]
        if any(abs((g or 0) - w) > 0.01 for g, w in zip(row, want)): bad += 1; print(name, k, row, want) if bad < 5 else None
    tot = [sum(exp[k][j] for k in keys) for j in range(5)]
    print(f'{name}: {len(keys)} rows · month {tot[0]:,.0f} · prev {tot[1]:,.0f} · YTD {tot[2]:,.0f} · ITD {tot[3]:,.0f} · pending {tot[4]:,.0f}')
print('DIV (live) blanks:', sum(1 for r in dt.iter_rows(min_row=2, values_only=True) if r[ix['Statement key']] and not r[ix['DIV (live)']]))
sys.exit(1 if bad else 0)
